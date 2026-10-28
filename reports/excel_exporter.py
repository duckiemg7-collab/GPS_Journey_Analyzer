from __future__ import annotations

from datetime import datetime
from copy import copy
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell import WriteOnlyCell

from config.settings import ALGORITHM_VERSION
from core.display_labels import event_type_label, warning_status_label
from core.models import AnalysisResult


NAVY = "17324D"
BLUE = "2F6B9A"
LIGHT_BLUE = "DCEAF5"
GREEN = "2E8B57"
AMBER = "F4B942"
RED = "C0392B"
WHITE = "FFFFFF"
TEXT = "243746"


# Chỉ dùng cho lớp trình bày trong Excel. Tên cột nội bộ được giữ nguyên để
# thuật toán, GUI và các bước đối soát không bị ảnh hưởng.
EXCEL_COLUMN_LABELS = {
    "detail_index": "STT nội bộ",
    "date": "Ngày",
    "device_id": "Thiết bị",
    "source_file": "File nguồn",
    "source_sheet": "Sheet nguồn",
    "source_row": "Dòng nguồn",
    "timestamp": "Thời gian",
    "latitude": "Vĩ độ",
    "longitude": "Kinh độ",
    "device_status": "Trạng thái thiết bị",
    "device_speed_kmh": "Tốc độ thiết bị (km/h)",
    "gps_speed_kmh": "Tốc độ GPS (km/h)",
    "counter_speed_kmh": "Tốc độ theo công tơ (km/h)",
    "device_km": "Số km thiết bị",
    "device_km_delta": "Mức tăng km thiết bị",
    "device_km_reset": "Thiết bị reset công tơ",
    "device_original_km": "Km thiết bị",
    "device_unallocated_km": "Km thiết bị chưa phân bổ",
    "unallocated_device_km": "Km chưa phân bổ qua ngày",
    "device_corrected_km": "Km thiết bị sau xác minh",
    "confirmed_false_km": "Km sai đã xác nhận",
    "device_suspected_km": "Km thiết bị nghi vấn",
    "deducted_km": "Km đã trừ có căn cứ",
    "raw_distance_km": "Khoảng cách GPS thô (km)",
    "gps_raw_km": "Km GPS thô",
    "gps_filtered_km": "Km GPS lọc",
    "gps_unverified_km": "Km GPS chưa xác minh",
    "gps_confirmed_outlier_km": "Km nhiễu GPS đã loại",
    "gps_filtered_segment_km": "Km đoạn GPS được giữ",
    "gps_unverified_segment_km": "Km đoạn GPS chưa xác minh",
    "gps_confirmed_outlier_segment_km": "Km đoạn nhiễu GPS đã loại",
    "gps_segment_raw_km": "Km GPS thô của đoạn",
    "gps_segment_reference_km": "Km tham khảo của đoạn",
    "actual_unverified_km": "Km thực tế chưa xác minh",
    "device_counter_span_km": "Km công tơ trong đoạn",
    "device_allocated_segment_km": "Km thiết bị đã phân bổ cho đoạn",
    "unallocated_device_segment_km": "Km thiết bị chưa phân bổ của đoạn",
    "delta_time_seconds": "Khoảng thời gian (giây)",
    "gps_segment_seconds": "Thời lượng đoạn GPS (giây)",
    "algorithm_status": "Kết quả xử lý",
    "warning_reason": "Lý do cảnh báo",
    "event_type": "Loại sự kiện",
    "device_anomaly_status": "Trạng thái bất thường km thiết bị",
    "confirmation_status": "Trạng thái xác minh",
    "evidence": "Căn cứ đánh giá",
    "gps_segment_class": "Phân loại đoạn GPS",
    "gps_segment_reason": "Lý do phân loại đoạn GPS",
    "gps_unverified_reason": "Lý do chưa xác minh",
    "reference_kind": "Loại số liệu tham khảo",
    "exclude_from_gps": "Loại khỏi km GPS",
    "no_fix_coordinate": "Không có tọa độ hợp lệ",
    "is_duplicate": "Bản ghi trùng",
    "invalid_timestamp": "Thời gian không hợp lệ",
    "timestamp_conflict": "Trùng thời gian khác dữ liệu",
    "invalid_coordinate": "Tọa độ không hợp lệ",
    "signal_loss_observed": "Ghi nhận mất tín hiệu",
    "counter_eligible": "Đủ điều kiện tính công tơ",
    "cross_day": "Đoạn đi qua ngày",
    "input_order": "Thứ tự dữ liệu đầu vào",
    "pair_start_index": "Chỉ số điểm đầu cặp",
    "gps_unverified_start_index": "Chỉ số đầu đoạn chưa xác minh",
    "gps_unverified_end_index": "Chỉ số cuối đoạn chưa xác minh",
    "gps_segment_start_index": "Chỉ số đầu đoạn GPS",
    "gps_segment_end_index": "Chỉ số cuối đoạn GPS",
    "warning_event_id": "Mã nhóm cảnh báo",
    "segment_id": "Mã đoạn GPS",
    "counter_edge_id": "Mã đoạn công tơ",
    "start_timestamp": "Thời gian bắt đầu",
    "end_timestamp": "Thời gian kết thúc",
    "start_latitude": "Vĩ độ bắt đầu",
    "start_longitude": "Kinh độ bắt đầu",
    "end_latitude": "Vĩ độ kết thúc",
    "end_longitude": "Kinh độ kết thúc",
    "start_source_file": "File nguồn điểm đầu",
    "start_source_sheet": "Sheet nguồn điểm đầu",
    "start_source_row": "Dòng nguồn điểm đầu",
    "end_source_file": "File nguồn điểm cuối",
    "end_source_sheet": "Sheet nguồn điểm cuối",
    "end_source_row": "Dòng nguồn điểm cuối",
    "start_device_km": "Số km thiết bị tại điểm đầu",
    "end_device_km": "Số km thiết bị tại điểm cuối",
    "start_device_status": "Trạng thái thiết bị tại điểm đầu",
    "end_device_status": "Trạng thái thiết bị tại điểm cuối",
    "start_device_speed_kmh": "Tốc độ thiết bị tại điểm đầu (km/h)",
    "end_device_speed_kmh": "Tốc độ thiết bị tại điểm cuối (km/h)",
    "warning_count": "Số cảnh báo",
    "warnings": "Chi tiết cảnh báo",
    "record_count": "Số bản ghi",
    "valid_record_count": "Số bản ghi hợp lệ",
    "duplicate_record_count": "Số bản ghi trùng",
    "excluded_record_count": "Số bản ghi bị loại",
    "gap_count": "Số đoạn mất tín hiệu",
    "unique_warning_events": "Số nhóm cảnh báo",
    "unverified_unknown_reference_count": "Số đoạn chưa đủ căn cứ",
    "device_period_original_km": "Km thiết bị toàn kỳ",
    "device_daily_allocated_km": "Km thiết bị đã phân bổ theo ngày",
    "counter_reconciliation_difference_km": "Chênh lệch đối soát công tơ (km)",
    "device_period_corrected_km": "Km thiết bị toàn kỳ sau xác minh",
    "raw_records": "Tổng bản ghi nguồn",
    "valid_records": "Bản ghi hợp lệ",
    "duplicate_records": "Bản ghi trùng",
    "excluded_records": "Bản ghi bị loại",
    "invalid_timestamp_records": "Bản ghi sai thời gian",
    "warning_records": "Bản ghi có cảnh báo",
    "gap_segments": "Số đoạn GPS chưa xác minh",
    "signal_loss_segments": "Số đoạn mất tín hiệu",
    "counter_resets": "Số lần reset công tơ",
    "difference_device_vs_gps_km": "Chênh lệch thiết bị và GPS (km)",
    "label": "Ngày / thiết bị",
    "gps_filtered_km_before": "Km GPS lọc trước sửa",
    "gps_filtered_km_after": "Km GPS lọc sau sửa",
    "confirmed_false_km_before": "Km sai xác nhận trước sửa",
    "confirmed_false_km_after": "Km sai xác nhận sau sửa",
    "device_corrected_km_before": "Km thiết bị hiệu chỉnh trước sửa",
    "device_corrected_km_after": "Km thiết bị hiệu chỉnh sau sửa",
    "gps_change_km": "Mức thay đổi km GPS",
    "change_reason": "Lý do thay đổi",
    "algorithm_status_before": "Kết quả xử lý trước sửa",
    "algorithm_status_after": "Kết quả xử lý sau sửa",
    "gps_minus_counter_km": "GPS trừ công tơ (km)",
    "gps_minus_period_counter_km": "GPS trừ công tơ toàn kỳ (km)",
    "trusted_gps_at_reported_stop_no_counter_increase_km": "Km GPS khi báo dừng và công tơ không tăng",
    "other_trusted_gps_minus_counter_km": "Chênh lệch GPS tin cậy khác so với công tơ",
    "counter_km_outside_trusted_gps_spans": "Km công tơ ngoài các đoạn GPS tin cậy",
}

BOOLEAN_COLUMNS = {
    "exclude_from_gps", "no_fix_coordinate", "is_duplicate", "invalid_timestamp",
    "timestamp_conflict", "invalid_coordinate", "signal_loss_observed", "counter_eligible",
    "device_km_reset", "cross_day",
}

SEGMENT_CLASS_LABELS = {
    "TRUSTED": "Đáng tin cậy",
    "UNVERIFIED": "Chưa xác minh",
    "OUTLIER": "Nhiễu GPS đã loại",
    "CONFIRMED_OUTLIER": "Nhiễu GPS đã xác nhận",
}

CONFIRMATION_LABELS = {
    "CONFIRMED": "Đã xác nhận",
    "NEED_REVIEW": "Cần kiểm tra",
}

MODE_LABELS = {
    "Conservative": "Thận trọng",
    "Balanced": "Cân bằng",
    "Aggressive": "Mạnh",
}

GPS_REASON_LABELS = {
    "GPS day boundary": "Đoạn GPS đi qua ranh giới ngày",
    "Signal loss": "Mất tín hiệu GPS",
    "GPS gap": "Khoảng cách thời gian GPS quá lớn",
    "Stationary drift": "Trôi tọa độ khi xe dừng",
    "GPS teleport": "GPS nhảy vị trí phi thực tế",
    "Confirmed bad GPS geometry": "Hình học GPS sai đã được xác nhận",
    "Missing coordinate fix; reference unavailable": "Thiếu tọa độ hợp lệ; không có khoảng cách tham khảo",
    "Signal/time/day/coordinate gap; straight-line reference only": "Gián đoạn tín hiệu, thời gian, ngày hoặc tọa độ; chỉ có khoảng cách đường thẳng tham khảo",
    "Unresolved GPS outlier; straight-line reference only": "Điểm GPS bất thường chưa xác minh; chỉ có khoảng cách đường thẳng tham khảo",
    "Rebuilt accepted endpoints": "Tái dựng từ các điểm đầu cuối được chấp nhận",
    "Continuous accepted endpoints": "Các điểm đầu cuối hợp lệ liên tục",
    "Relocated GPS; incoming route unknown": "GPS đổi vị trí; chưa biết tuyến đường đi tới",
    "Continuous samples after GPS reacquisition": "Mẫu liên tục sau khi GPS bắt lại tín hiệu",
    "Reacquired GPS; intervening route unknown": "GPS bắt lại tín hiệu; chưa biết tuyến đường trong khoảng gián đoạn",
}


def _warning_summary_label(value: object) -> object:
    if not isinstance(value, str) or not value:
        return value
    parts = []
    for item in value.split(","):
        label, separator, count = item.strip().partition(":")
        parts.append(f"{warning_status_label(label)}: {count.strip()}" if separator else warning_status_label(label))
    return ", ".join(parts)


def _localize_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Tạo bản sao dành cho Excel, không thay đổi DataFrame nội bộ."""
    localized = frame.copy()
    status_columns = {"algorithm_status", "device_anomaly_status", "algorithm_status_before", "algorithm_status_after"}
    for column in status_columns.intersection(localized.columns):
        localized[column] = localized[column].map(
            lambda value: warning_status_label(value) if pd.notna(value) else value
        )
    if "event_type" in localized:
        localized["event_type"] = localized["event_type"].map(
            lambda value: event_type_label(value) if pd.notna(value) else value
        )
    if "warnings" in localized:
        localized["warnings"] = localized["warnings"].map(_warning_summary_label)
    if "gps_segment_class" in localized:
        localized["gps_segment_class"] = localized["gps_segment_class"].replace(SEGMENT_CLASS_LABELS)
    for column in {"gps_segment_reason", "gps_unverified_reason"}.intersection(localized.columns):
        localized[column] = localized[column].replace(GPS_REASON_LABELS)
    if "confirmation_status" in localized:
        localized["confirmation_status"] = localized["confirmation_status"].replace(CONFIRMATION_LABELS)
    if "reference_kind" in localized:
        localized["reference_kind"] = localized["reference_kind"].replace({
            "Haversine endpoint straight line; actual route unknown for UNVERIFIED":
                "Đường thẳng Haversine giữa hai đầu mút; chưa biết hành trình thực tế",
        })
    if "evidence" in localized:
        localized["evidence"] = localized["evidence"].replace({
            "GPS/counter inconsistency; delayed upload, movement and reset cannot be ruled out. No exact false amount proven.":
                "GPS và công tơ không nhất quán; chưa loại trừ gửi dữ liệu trễ, xe di chuyển hoặc reset công tơ. Chưa chứng minh được chính xác số km sai.",
        })
    for column in BOOLEAN_COLUMNS.intersection(localized.columns):
        localized[column] = localized[column].map(
            lambda value: "Có" if value is True or isinstance(value, np.bool_) and bool(value)
            else "Không" if value is False or isinstance(value, np.bool_) and not bool(value)
            else value
        )
    return localized.rename(columns=EXCEL_COLUMN_LABELS)


def _safe(value):
    if value is pd.NA or value is pd.NaT or (isinstance(value, (float, np.floating)) and np.isnan(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.generic):
        return value.item()
    return value


def _write_dataframe(ws, frame: pd.DataFrame, start_row: int = 1) -> tuple[int, int]:
    for col, name in enumerate(frame.columns, 1):
        cell = ws.cell(start_row, col, str(name))
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, row in enumerate(frame.itertuples(index=False, name=None), start_row + 1):
        for c, value in enumerate(row, 1):
            safe = _safe(value)
            ws.cell(r, c, "Không đủ căn cứ" if safe is None else safe)
    last_row = start_row + len(frame)
    last_col = len(frame.columns)
    if last_col:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(last_col)}{max(start_row, last_row)}"
    return last_row, last_col


def _finish_sheet(ws, freeze: str = "A2", max_width: int = 42) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = BLUE
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        sample = list(column_cells[:250])
        width = max((len(str(c.value)) for c in sample if c.value is not None), default=8) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), max_width)
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name, font.sz = "Arial", 10
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.vertical = "center"
            cell.alignment = alignment


def _format_detail(ws, headers: list[str], last_row: int) -> None:
    if last_row < 2:
        return
    for idx, name in enumerate(headers, 1):
        if name == "timestamp" or name.endswith("_timestamp"):
            for cell in ws.iter_cols(idx, idx, 2, last_row):
                cell[0].number_format = "dd/mm/yyyy hh:mm:ss"
        if name.endswith("_km") or name in {"raw_distance_km", "gps_filtered_segment_km"}:
            for cell in ws.iter_cols(idx, idx, 2, last_row):
                cell[0].number_format = "0.000"
        if name.endswith("_kmh"):
            for cell in ws.iter_cols(idx, idx, 2, last_row):
                cell[0].number_format = "0.0"


def reconciliation_tables(result: AnalysisResult) -> dict[str, pd.DataFrame]:
    evidence = result.counter_evidence
    segment_columns = [
        "segment_id", "device_id", "gps_segment_class", "gps_segment_reason",
        "start_timestamp", "end_timestamp", "start_latitude", "start_longitude",
        "end_latitude", "end_longitude", "start_source_file", "start_source_sheet",
        "start_source_row", "end_source_file", "end_source_sheet", "end_source_row",
        "gps_segment_start_index", "gps_segment_end_index", "gps_segment_seconds",
        "gps_segment_reference_km", "actual_unverified_km", "reference_kind",
    ]
    counter_columns = [
        "counter_edge_id", "device_id", "start_timestamp", "end_timestamp",
        "start_latitude", "start_longitude", "end_latitude", "end_longitude",
        "start_source_file", "start_source_row", "end_source_file", "end_source_row",
        "start_device_km", "end_device_km", "device_km_delta", "raw_distance_km",
        "delta_time_seconds", "gps_speed_kmh", "counter_speed_kmh", "device_suspected_km",
        "deducted_km", "confirmation_status", "evidence",
    ]
    segments = result.segments
    if segments.empty:
        unverified = pd.DataFrame(columns=segment_columns)
    else:
        unverified = segments.loc[segments.gps_segment_class.eq("UNVERIFIED")].reindex(columns=segment_columns)
    if evidence.empty:
        suspected = confirmed = pd.DataFrame(columns=counter_columns)
    else:
        suspected = evidence.loc[evidence.confirmation_status.eq("NEED_REVIEW")].reindex(columns=counter_columns)
        confirmed = evidence.loc[evidence.confirmation_status.eq("CONFIRMED")].reindex(columns=counter_columns)
    return {
        "Km thiết bị nghi cộng sai": suspected,
        "Km thiết bị cộng sai xác nhận": confirmed,
        "Các đoạn GPS chưa xác minh": unverified,
        "Đối soát toàn kỳ": result.reconciliation,
    }


def export_reconciliation_report(
    result: AnalysisResult, output_path: str | Path,
    comparison: pd.DataFrame | None = None, per_file: pd.DataFrame | None = None,
    extra_tables: dict[str, pd.DataFrame] | None = None,
) -> Path:
    """Stream the full sample audit without retaining millions of Excel cells."""
    path = Path(output_path).resolve()
    if path in {p.resolve() for p in result.files}:
        raise ValueError("Không được ghi đè file Excel nguồn")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    detail_columns = [
        "device_id", "source_file", "source_sheet", "source_row", "timestamp", "latitude", "longitude",
        "device_status", "device_speed_kmh", "device_km", "device_km_delta", "raw_distance_km",
        "delta_time_seconds", "gps_speed_kmh", "gps_filtered_segment_km", "gps_unverified_segment_km",
        "gps_confirmed_outlier_segment_km", "confirmed_false_km", "device_anomaly_status",
        "algorithm_status", "exclude_from_gps", "warning_event_id", "warning_reason",
        "is_duplicate", "invalid_timestamp", "timestamp_conflict", "no_fix_coordinate", "signal_loss_observed",
        "gps_segment_start_index", "gps_segment_end_index", "gps_segment_class", "gps_segment_reason",
    ]
    tables = {
        "Tổng hợp theo ngày": result.daily,
        "Chi tiết GPS": result.details.reindex(columns=detail_columns).reset_index(names="detail_index"),
        "Các sự kiện nhiễu": result.events.reindex(columns=detail_columns).reset_index(names="detail_index"),
        **reconciliation_tables(result),
    }
    if comparison is not None:
        tables["So sánh trước sau"] = comparison
    if per_file is not None:
        tables["Tổng hợp theo file"] = per_file
    tables.update(extra_tables or {})
    log = [
        ("Phiên bản thuật toán", ALGORITHM_VERSION),
        ("Chế độ lọc", MODE_LABELS.get(result.settings.mode, result.settings.mode)),
        ("Tốc độ GPS tối đa (km/h)", result.settings.max_plausible_speed_kmh),
        ("Ngưỡng xe dừng (km/h)", result.settings.stop_speed_kmh),
        ("Bán kính dừng (m)", result.settings.stop_radius_meters),
        ("Khoảng mất GPS tối đa (giây)", result.settings.max_gps_gap_seconds),
    ]
    log += [("Định nghĩa", text) for text in [
        "Km GPS chưa xác minh là khoảng cách Haversine giữa hai đầu mút, không phải quãng đường thực tế chưa biết.",
        "Km nhiễu GPS đã loại là hình học GPS thô bị loại, không cộng thêm vào các nhóm quãng đường tái dựng.",
        "Mỗi khoảng GPS thô hợp lệ chỉ thuộc đúng một đoạn GPS tái dựng.",
        "Tổng, hợp lệ, trùng, bị loại và cảnh báo là số bản ghi; các nhóm bị loại, trùng và cảnh báo có thể giao nhau.",
        "Một nhóm cảnh báo gồm các bản ghi liên tiếp cùng loại, thiết bị và ngày, cách nhau không quá 120 giây.",
        "Km thiết bị chưa phân bổ là mức tăng qua ranh giới ngày; ngày của điểm cuối chỉ dùng làm nhóm kiểm tra.",
        "Km thực tế chưa xác minh luôn hiển thị 'Không đủ căn cứ', không thay bằng 0.",
        "GPS nhảy vị trí không tự chứng minh công tơ sai; các trường hợp hiện tại vẫn cần kiểm tra.",
    ]]
    log += [("File nguồn", str(p)) for p in result.files]
    log += [("Nhật ký", message) for message in result.logs]
    tables["Nhật ký thuật toán và cấu hình"] = pd.DataFrame(log, columns=["Mục", "Nội dung"])
    for name, original_frame in tables.items():
        frame = _localize_frame(original_frame)
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for idx, column in enumerate(original_frame.columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = 23 if "timestamp" in str(column) else 18
        headers = []
        for column in frame.columns:
            cell = WriteOnlyCell(ws, str(column))
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Arial", bold=True, color=WHITE)
            headers.append(cell)
        ws.append(headers)
        for values in frame.itertuples(index=False, name=None):
            cells = []
            for column, value in zip(original_frame.columns, values):
                value = _safe(value)
                if value is None:
                    value = "Không đủ căn cứ"
                cell = WriteOnlyCell(ws, value)
                if "timestamp" in str(column) or column == "date":
                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                elif str(column).endswith("_km") or str(column).endswith("_kmh"):
                    cell.number_format = "0.000"
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(frame.columns)))}{len(frame) + 1}"
    wb.save(path)
    return path


def export_report(result: AnalysisResult, output_path: str | Path) -> Path:
    path = Path(output_path).resolve()
    if path in {p.resolve() for p in result.files}:
        raise ValueError("Không được ghi đè file Excel nguồn")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    summary = result.daily[[
        "date", "device_id", "device_original_km", "gps_filtered_km",
        "device_corrected_km", "confirmed_false_km", "warning_count", "warnings", "record_count",
    ]].copy()
    summary = _localize_frame(summary)
    ws = wb.create_sheet("Tổng hợp theo ngày")
    last_row, _ = _write_dataframe(ws, summary)
    total_row = last_row + 2
    ws.cell(total_row, 1, "TỔNG TOÀN KỲ")
    ws.cell(total_row, 1).font = Font(name="Arial", bold=True, color=WHITE)
    ws.cell(total_row, 1).fill = PatternFill("solid", fgColor=BLUE)
    for col in (3, 4, 5, 6, 7, 9):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}2:{letter}{last_row})")
        ws.cell(total_row, col).font = Font(name="Arial", bold=True)
        if col in (3, 4, 5, 6):
            ws.cell(total_row, col).number_format = "0.000"
    for row in range(2, last_row + 1):
        ws.cell(row, 1).number_format = "dd/mm/yyyy"
        for col in (3, 4, 5, 6):
            ws.cell(row, col).number_format = "0.000"
    ws.conditional_formatting.add(
        f"G2:G{last_row}",
        FormulaRule(formula=["G2>0"], fill=PatternFill("solid", fgColor="FFF1CC")),
    )
    _finish_sheet(ws)

    event_columns = [
        "timestamp",
        "device_id",
        "latitude",
        "longitude",
        "raw_distance_km",
        "delta_time_seconds",
        "gps_speed_kmh",
        "device_speed_kmh",
        "device_km",
        "device_km_delta",
        "confirmed_false_km",
        "event_type",
        "warning_reason",
        "algorithm_status",
        "source_file",
        "source_sheet",
        "source_row",
    ]
    event_columns = [c for c in event_columns if c in result.events.columns]
    localized_events = _localize_frame(result.events[event_columns])
    ws_events = wb.create_sheet("Điểm GPS bất thường")
    event_last, _ = _write_dataframe(ws_events, localized_events)
    _format_detail(ws_events, event_columns, event_last)
    _finish_sheet(ws_events, "A2", 50)

    for name, frame in reconciliation_tables(result).items():
        original_columns = list(frame.columns)
        frame = _localize_frame(frame)
        audit_ws = wb.create_sheet(name)
        last, _ = _write_dataframe(audit_ws, frame)
        _format_detail(audit_ws, original_columns, last)
        _finish_sheet(audit_ws, "A2", 45)

    comparison = result.daily[
        [
            "date",
            "device_id",
            "device_original_km",
            "gps_filtered_km",
            "device_corrected_km",
            "confirmed_false_km",
        ]
    ].copy()
    comparison["difference_device_vs_gps_km"] = (
        comparison["device_original_km"] - comparison["gps_filtered_km"]
    )
    comparison_columns = list(comparison.columns)
    comparison = _localize_frame(comparison)
    ws_compare = wb.create_sheet("Đối chiếu km")
    comp_last, _ = _write_dataframe(ws_compare, comparison)
    _format_detail(ws_compare, comparison_columns, comp_last)
    _finish_sheet(ws_compare)

    ws_chart = wb.create_sheet("Biểu đồ và thống kê")
    ws_chart.sheet_view.showGridLines = False
    ws_chart["A2"] = "GPS Journey Analyzer"
    ws_chart["A2"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    ws_chart["A4"] = "Chỉ số"
    ws_chart["B4"] = "Giá trị"
    totals = result.totals
    metrics = [
        ("Tổng số ngày", totals["days"]),
        ("Tổng bản ghi", totals["records"]),
        ("Km thiết bị báo", totals["device_original_km"]),
        ("Km GPS sau lọc", totals["gps_filtered_km"]),
        ("Km thiết bị hiệu chỉnh", totals["device_corrected_km"]),
        ("Km sai xác nhận", totals["confirmed_false_km"]),
        ("Bản ghi có cảnh báo", totals["events"]),
        ("Nhóm sự kiện cảnh báo", totals.get("unique_warning_events")),
    ]
    for col in (1, 2):
        ws_chart.cell(4, col).fill = PatternFill("solid", fgColor=NAVY)
        ws_chart.cell(4, col).font = Font(name="Arial", bold=True, color=WHITE)
    for row, (label, value) in enumerate(metrics, 5):
        ws_chart.cell(row, 1, label)
        ws_chart.cell(row, 2, _safe(value))
        if "Km" in label:
            ws_chart.cell(row, 2).number_format = "0.000"
    chart_start = 15
    chart_data = result.daily[result.daily["date"].notna()].copy()
    chart_data["label"] = chart_data.apply(
        lambda r: f"{pd.Timestamp(r['date']).strftime('%d/%m')} {r['device_id']}", axis=1
    )
    chart_table = chart_data[
        ["label", "device_original_km", "gps_filtered_km", "device_corrected_km"]
    ]
    chart_table = _localize_frame(chart_table)
    _write_dataframe(ws_chart, chart_table, chart_start)
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "So sánh quãng đường theo ngày (km)"
    chart.y_axis.title = "km"
    chart.x_axis.title = "Ngày / thiết bị"
    chart.add_data(
        Reference(ws_chart, min_col=2, max_col=4, min_row=chart_start, max_row=chart_start + len(chart_table)),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(ws_chart, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(chart_table))
    )
    chart.height = 9
    chart.width = 20
    ws_chart.add_chart(chart, "D4")
    _finish_sheet(ws_chart, "A1", 28)

    ws_log = wb.create_sheet("Nhật ký xử lý")
    log_rows = [
        ("Thời điểm xuất", datetime.now()),
        ("Phiên bản thuật toán", ALGORITHM_VERSION),
        ("Chế độ lọc", MODE_LABELS.get(result.settings.mode, result.settings.mode)),
        ("Tốc độ GPS tối đa (km/h)", result.settings.max_plausible_speed_kmh),
        ("Ngưỡng xe dừng (km/h)", result.settings.stop_speed_kmh),
        ("Bán kính dừng (m)", result.settings.stop_radius_meters),
        ("Khoảng mất GPS tối đa (giây)", result.settings.max_gps_gap_seconds),
        ("Giới hạn", "GPS lọc là độ dài tuyến nối điểm được chấp nhận, không phải quãng đường thực tuyệt đối."),
        ("Giới hạn", "Đoạn mất GPS không được tự nối; kết quả có thể thấp hơn thực tế."),
        ("Km chưa xác minh", "Chỉ là Haversine đường thẳng tham khảo; quãng đường thực chưa biết. Không cộng với GPS lọc."),
        ("GPS outlier loại", "Độ dài hình học thô đã loại; không cộng vào các nhóm quãng đường tái dựng."),
        ("Km chưa phân bổ", "Mức tăng công tơ qua ngày giữ riêng; tổng ngày + chưa phân bổ = tổng kỳ độc lập."),
        ("Thống kê", "Raw gồm mọi dòng; valid = không cảnh báo; duplicate và excluded có thể giao với warning. Sự kiện là nhóm cảnh báo liên tiếp cùng loại/xe/ngày, cách tối đa 120 giây."),
    ]
    for file in result.files:
        log_rows.append(("File nguồn", str(file)))
    for message in result.logs:
        log_rows.append(("Nhật ký", message))
    _write_dataframe(ws_log, pd.DataFrame(log_rows, columns=["Mục", "Nội dung"]))
    ws_log["B2"].number_format = "dd/mm/yyyy hh:mm:ss"
    _finish_sheet(ws_log, "A2", 100)

    wb.save(path)
    # Kiểm tra file có thể mở lại và đủ sheet sau khi ghi.
    check = load_workbook(path, read_only=True, data_only=False)
    expected = {
        "Tổng hợp theo ngày",
        "Điểm GPS bất thường",
        "Đối chiếu km",
        "Biểu đồ và thống kê",
        "Nhật ký xử lý",
    }
    missing = expected.difference(check.sheetnames)
    check.close()
    if missing:
        raise RuntimeError(f"Báo cáo thiếu sheet: {sorted(missing)}")
    return path
