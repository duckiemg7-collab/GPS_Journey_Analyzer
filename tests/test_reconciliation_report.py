from pathlib import Path

from openpyxl import load_workbook
import pytest

from core.pipeline import analyze_records
from reports.excel_exporter import export_report, export_reconciliation_report
from tests.conftest import make_frame


def result_for_test():
    raw = make_frame([
        ("2026-10-01 23:59:50", 10, 106, 30, 10),
        ("2026-10-02 00:00:10", 10.001, 106, 30, 10.1),
        ("2026-10-02 00:01:10", 10.002, 106, 30, 10.2),
    ])
    return analyze_records(raw, [Path("synthetic.xlsx")])


def test_existing_report_preserves_summary_columns_and_adds_evidence(tmp_path):
    result = result_for_test()
    path = export_report(result, tmp_path / "report.xlsx")
    wb = load_workbook(path, read_only=True)
    ws = wb["Tổng hợp theo ngày"]
    assert [ws.cell(1, i).value for i in range(1, 10)] == [
        "Ngày", "Thiết bị", "Km thiết bị", "Km GPS lọc",
        "Km thiết bị sau xác minh", "Km sai đã xác nhận", "Số cảnh báo",
        "Chi tiết cảnh báo", "Số bản ghi",
    ]
    assert ws.cell(5, 3).value == "=SUM(C2:C3)"
    assert ws.cell(5, 9).value == "=SUM(I2:I3)"
    assert ws.cell(3, 8).value == "Mất tín hiệu GPS: 1"
    for name in ("Km thiết bị nghi cộng sai", "Km thiết bị cộng sai xác nhận", "Các đoạn GPS chưa xác minh", "Đối soát toàn kỳ"):
        assert name in wb.sheetnames
    assert "Chi tiết GPS" not in wb.sheetnames
    events = wb["Điểm GPS bất thường"]
    headers = [cell.value for cell in next(events.iter_rows())]
    assert "Kết quả xử lý" in headers
    assert "Khoảng cách GPS thô (km)" in headers
    assert "algorithm_status" not in headers
    status_column = headers.index("Kết quả xử lý") + 1
    statuses = {events.cell(row, status_column).value for row in range(2, events.max_row + 1)}
    assert statuses == {"Mất tín hiệu GPS"}
    wb.close()


def test_streamed_report_has_true_endpoints_and_unknown_actual_distance(tmp_path):
    path = export_reconciliation_report(result_for_test(), tmp_path / "audit.xlsx")
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb["Các đoạn GPS chưa xác minh"]
    rows = list(ws.values)
    record = dict(zip(rows[0], rows[1]))
    assert record["Thời gian bắt đầu"].isoformat() == "2026-10-01T23:59:50"
    assert record["Thời gian kết thúc"].isoformat() == "2026-10-02T00:00:10"
    assert record["Km thực tế chưa xác minh"] == "Không đủ căn cứ"
    assert record["Dòng nguồn điểm đầu"] == 4
    assert record["Dòng nguồn điểm cuối"] == 5
    assert record["Phân loại đoạn GPS"] == "Chưa xác minh"
    assert "start_timestamp" not in rows[0]
    wb.close()


@pytest.mark.parametrize("exporter", [export_report, export_reconciliation_report])
def test_export_cannot_overwrite_source(tmp_path, exporter):
    result = result_for_test()
    source = tmp_path / "source.xlsx"
    result.files = [source]
    with pytest.raises(ValueError, match="nguồn"):
        exporter(result, source)
    assert not source.exists()
