from __future__ import annotations

import re
import tempfile
import unicodedata
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import load_workbook


class ExcelReadError(RuntimeError):
    pass


def _plain(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").lower())
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn").replace("đ", "d").strip()


def _repair_invalid_styles(source: Path, target: Path) -> None:
    """Sửa mã màu 5/6 ký tự trên bản sao; không thay đổi file nguồn."""
    with ZipFile(source) as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/styles.xml":
                xml = data.decode("utf-8")

                def fix(match: re.Match[str]) -> str:
                    value = match.group(1)
                    if len(value) == 6:
                        value = "FF" + value
                    elif len(value) < 8:
                        value = "FF" + value.zfill(6)
                    return f'rgb="{value}"'

                data = re.sub(r'rgb="([0-9A-Fa-f]+)"', fix, xml).encode("utf-8")
            dst.writestr(info, data)


def _openable_copy(path: Path, temp_dir: Path, logs: list[str]) -> Path:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        wb.close()
        return path
    except ValueError as exc:
        if "stylesheet" not in str(exc).lower() and "colors must" not in str(exc).lower():
            raise
        repaired = temp_dir / path.name
        _repair_invalid_styles(path, repaired)
        logs.append(f"Đã sửa XML style không chuẩn trên bản sao tạm: {path.name}")
        return repaired


def _find_header(path: Path, sheet_name: str) -> int:
    preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=40, dtype=str)
    required_groups = ({"thoi diem", "timestamp", "time"}, {"toa do", "coordinate"})
    for idx, row in preview.iterrows():
        values = {_plain(v) for v in row.dropna()}
        if all(any(alias in values for alias in group) for group in required_groups):
            return int(idx)
    raise ExcelReadError(f"Không tìm thấy dòng tiêu đề trong sheet {sheet_name!r}")


def _pick(columns: list[str], aliases: tuple[str, ...], required: bool = True) -> str | None:
    normalized = {_plain(col): col for col in columns}
    for alias in aliases:
        if alias in normalized:
            return normalized[alias]
    if required:
        raise ExcelReadError(f"Thiếu cột bắt buộc: {aliases[0]}")
    return None


def _device_id(path: Path, sheet: str, first_cells: list[object]) -> str:
    haystack = " ".join(str(x or "") for x in first_cells) + " " + path.stem + " " + sheet
    match = re.search(r"(?:XE\s+)?([0-9]{2}[A-Z]{1,3}[0-9A-Z]{4,})", haystack, re.I)
    return match.group(1).upper() if match else sheet.strip()


def _parse_coordinates(series: pd.Series) -> tuple[pd.Series, pd.Series]:
    extracted = series.astype("string").str.extract(
        r"^\s*([+-]?\d+(?:[.,]\d+)?)([-+]\d+(?:[.,]\d+)?)\s*$"
    )
    lat = pd.to_numeric(extracted[0].str.replace(",", ".", regex=False), errors="coerce")
    lon = pd.to_numeric(extracted[1].str.replace(",", ".", regex=False), errors="coerce")
    return lat, lon


def _parse_timestamps(series: pd.Series) -> pd.Series:
    """Đọc timestamp dạng ``HH:MM:SS DD-MM-YYYY`` mà không suy đoán sai."""
    values = series.astype("string").str.strip()
    parsed = pd.to_datetime(
        values,
        format="%H:%M:%S %d-%m-%Y",
        errors="coerce",
    )
    iso = values.str.match(r"^\d{4}-\d{2}-\d{2}(?:[T ]|$)", na=False)
    parsed.loc[iso] = pd.to_datetime(values.loc[iso], format="ISO8601", errors="coerce")
    missing = parsed.isna() & values.notna()
    if missing.any():
        # Fallback cho các mẫu Excel khác (datetime thật, có giây thiếu,
        # hoặc định dạng ISO), nhưng vẫn yêu cầu day-first.
        parsed.loc[missing] = pd.to_datetime(
            values.loc[missing],
            format="mixed",
            dayfirst=True,
            errors="coerce",
        )
    return parsed


def discover_excel_files(paths: list[str | Path]) -> list[Path]:
    found: set[Path] = set()
    for raw in paths:
        path = Path(raw).expanduser().resolve()
        if path.is_dir():
            found.update(p for p in path.rglob("*.xlsx") if not p.name.startswith("~$"))
            found.update(p for p in path.rglob("*.xlsm") if not p.name.startswith("~$"))
        elif path.suffix.lower() in {".xlsx", ".xlsm"} and not path.name.startswith("~$"):
            found.add(path)
    return sorted(found)


def read_excel_files(paths: list[str | Path]) -> tuple[pd.DataFrame, list[Path], list[str]]:
    files = discover_excel_files(paths)
    if not files:
        raise ExcelReadError("Không tìm thấy file .xlsx/.xlsm hợp lệ.")
    frames: list[pd.DataFrame] = []
    logs: list[str] = []
    with tempfile.TemporaryDirectory(prefix="gps_ja_") as temp:
        temp_dir = Path(temp)
        for source in files:
            try:
                readable = _openable_copy(source, temp_dir, logs)
                wb = load_workbook(readable, read_only=True, data_only=True)
                sheets = wb.sheetnames
                wb.close()
                for sheet in sheets:
                    header = _find_header(readable, sheet)
                    frame = pd.read_excel(readable, sheet_name=sheet, header=header, dtype=str)
                    frame = frame.dropna(how="all").reset_index(drop=True)
                    if frame.empty:
                        continue
                    columns = [str(c).strip() for c in frame.columns]
                    frame.columns = columns
                    time_col = _pick(columns, ("thoi diem", "timestamp", "time"))
                    speed_col = _pick(columns, ("van toc", "speed"), False)
                    km_col = _pick(columns, ("so km", "km", "quang duong"), False)
                    coord_col = _pick(columns, ("toa do", "coordinate", "gps coordinate"))
                    status_col = _pick(columns, ("trang thai", "status"), False)
                    title = pd.read_excel(readable, sheet_name=sheet, header=None, nrows=2, dtype=str)
                    device = _device_id(source, sheet, title.fillna("").values.ravel().tolist())
                    lat, lon = _parse_coordinates(frame[coord_col])
                    normalized = frame.copy()
                    normalized.insert(0, "source_row", normalized.index + header + 2)
                    normalized.insert(0, "source_sheet", sheet)
                    normalized.insert(0, "source_file", source.name)
                    normalized.insert(0, "device_id", device)
                    normalized["timestamp"] = _parse_timestamps(frame[time_col])
                    normalized["latitude"] = lat
                    normalized["longitude"] = lon
                    normalized["device_speed_kmh"] = (
                        pd.to_numeric(frame[speed_col], errors="coerce") if speed_col else float("nan")
                    )
                    normalized["device_km"] = (
                        pd.to_numeric(frame[km_col], errors="coerce") if km_col else float("nan")
                    )
                    normalized["device_status"] = frame[status_col] if status_col else ""
                    frames.append(normalized)
                    logs.append(
                        f"Đọc {len(frame):,} dòng từ {source.name} / {sheet}; tiêu đề dòng {header + 1}."
                    )
            except Exception as exc:
                logs.append(f"LỖI {source.name}: {exc}")
    if not frames:
        raise ExcelReadError("Không đọc được dữ liệu từ các file đã chọn.\n" + "\n".join(logs))
    data = pd.concat(frames, ignore_index=True, sort=False)
    data["input_order"] = range(len(data))
    return data, files, logs
